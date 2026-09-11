"""The page on the left, its text on the right.

The point of this screen is comparison. OCR on a sefer is not going to be
perfect, so the useful thing is not a cleaner result, it is being able to see the
original and fix the text against it without opening two programs.

Clicking a paragraph in the text draws a box around the part of the page it came
from. That is what the per block coordinates are for, and it is what makes
correcting a two column page bearable.
"""

from __future__ import annotations

from pathlib import Path

from PySide6.QtCore import Qt, Signal
from PySide6.QtGui import QColor, QPainter, QPen, QPixmap
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QFileDialog,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QScrollArea,
    QSplitter,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from ..core.models import Document, PageText
from ..export import pdf as pdf_export
from ..export import txt as txt_export
from .widgets import Divider, button, caption

EXPORT_FILTERS = "Text (*.txt);;Word document (*.docx);;PDF (*.pdf)"


class PageImage(QScrollArea):
    """The original page, scaled to fit, with an optional highlight."""

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWidgetResizable(True)
        self.setAlignment(Qt.AlignCenter)
        self._label = QLabel("No page image")
        self._label.setAlignment(Qt.AlignCenter)
        self._label.setObjectName("Caption")
        self.setWidget(self._label)
        self._source: QPixmap | None = None
        self._page: PageText | None = None
        self._highlight: tuple[int, int, int, int] | None = None

    def show_page(self, page: PageText) -> None:
        self._page = page
        self._highlight = None
        path = Path(page.image_path or page.source_path)
        if not path.is_file():
            self._source = None
            self._label.setPixmap(QPixmap())
            self._label.setText(
                f"The page image is no longer at {path.name}.\n"
                "The text is unaffected."
            )
            return
        pixmap = QPixmap(str(path))
        self._source = None if pixmap.isNull() else pixmap
        if self._source is None:
            self._label.setText(f"{path.name} could not be shown as an image.")
        self._redraw()

    def highlight(self, bbox: tuple[int, int, int, int] | None) -> None:
        self._highlight = bbox
        self._redraw()

    def resizeEvent(self, event) -> None:      # noqa: N802
        super().resizeEvent(event)
        self._redraw()

    def _redraw(self) -> None:
        if self._source is None:
            return
        available = max(200, self.viewport().width() - 12)
        scaled = self._source.scaledToWidth(available, Qt.SmoothTransformation)

        if self._highlight and self._page and self._page.width:
            factor = scaled.width() / self._page.width
            x, y, w, h = self._highlight
            painter = QPainter(scaled)
            pen = QPen(QColor(29, 107, 87), max(2, int(2 * factor + 1)))
            painter.setPen(pen)
            painter.setBrush(QColor(29, 107, 87, 34))
            painter.drawRect(int(x * factor), int(y * factor),
                             int(w * factor), int(h * factor))
            painter.end()

        self._label.setPixmap(scaled)
        self._label.setText("")


class OcrReviewView(QWidget):
    """One document: page navigation, the image, and editable text."""

    def __init__(self, settings, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self._settings = settings
        self._document: Document | None = None
        self._page_index = 0
        self._block_starts: list[tuple[int, int]] = []   # (block index, char offset)
        self._loading = False

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(12)

        header = QHBoxLayout()
        header.setSpacing(12)
        titles = QVBoxLayout()
        titles.setSpacing(1)
        self.title = QLabel("Extracted text")
        self.title.setObjectName("PageTitle")
        self.subtitle = QLabel("")
        self.subtitle.setObjectName("Caption")
        titles.addWidget(self.title)
        titles.addWidget(self.subtitle)
        holder = QWidget()
        holder.setLayout(titles)
        header.addWidget(holder, 1)

        self.prev_button = button("Previous", lambda: self.go(self._page_index - 1))
        self.page_label = QLabel("")
        self.page_label.setObjectName("Mono")
        self.next_button = button("Next", lambda: self.go(self._page_index + 1))
        header.addWidget(self.prev_button)
        header.addWidget(self.page_label)
        header.addWidget(self.next_button)
        layout.addLayout(header)

        self.notes = caption("")
        self.notes.setVisible(False)
        layout.addWidget(self.notes)

        splitter = QSplitter(Qt.Horizontal)
        self.image = PageImage()
        splitter.addWidget(self.image)

        right = QWidget()
        right_layout = QVBoxLayout(right)
        right_layout.setContentsMargins(0, 0, 0, 0)
        right_layout.setSpacing(6)
        label = QLabel("EXTRACTED TEXT")
        label.setObjectName("SectionLabel")
        right_layout.addWidget(label)

        self.editor = QTextEdit()
        self.editor.setAcceptRichText(False)
        self.editor.cursorPositionChanged.connect(self._on_cursor)
        self.editor.textChanged.connect(self._on_edited)
        right_layout.addWidget(self.editor, 1)

        self.status = caption("")
        right_layout.addWidget(self.status)
        splitter.addWidget(right)
        splitter.setSizes([560, 560])
        layout.addWidget(splitter, 1)

        footer = QHBoxLayout()
        footer.setSpacing(8)
        footer.addWidget(button("Copy this page", self.copy_page))
        footer.addWidget(button("Copy everything", self.copy_all))
        footer.addStretch(1)
        footer.addWidget(button("Export", self.export, primary=True))
        layout.addLayout(footer)

    # -- loading ---------------------------------------------------------

    def load(self, document: Document) -> None:
        self._document = document
        self._page_index = 0
        name = Path(document.source_path).name or "Document"
        self.title.setText(name)
        # Deliberately not go(), which saves the editor into the current page
        # first. On a fresh load the editor still holds the previous document,
        # and saving that into the new one wiped every block it had just read.
        self._show(0)

    def go(self, index: int) -> None:
        """Move to another page, keeping whatever was typed on this one."""
        if not self._document or not self._document.pages:
            return
        self._save_edits()
        self._show(index)

    def _show(self, index: int) -> None:
        if not self._document or not self._document.pages:
            return
        self._page_index = max(0, min(index, len(self._document.pages) - 1))
        page = self.page

        self._loading = True
        self.image.show_page(page)
        text = "\n\n".join(
            b.text.strip() for b in sorted(page.blocks, key=lambda b: b.reading_order)
            if b.text.strip()
        )
        self.editor.setPlainText(text)
        self._loading = False

        # Where each block starts in the editor, so a click can find its box.
        self._block_starts = []
        offset = 0
        for block in sorted(page.blocks, key=lambda b: b.reading_order):
            if not block.text.strip():
                continue
            self._block_starts.append((page.blocks.index(block), offset))
            offset += len(block.text.strip()) + 2

        total = len(self._document.pages)
        self.page_label.setText(f"Page {self._page_index + 1} of {total}")
        self.prev_button.setEnabled(self._page_index > 0)
        self.next_button.setEnabled(self._page_index < total - 1)

        details = []
        if page.from_text_layer:
            details.append("read from the document's own text")
        else:
            details.append(f"{page.confidence:.0%} confident")
        if page.columns > 1:
            details.append(f"{page.columns} columns")
        if page.preprocessing:
            details.append(page.preprocessing.lower())
        self.subtitle.setText("  ·  ".join(details))

        notes = list(page.notes)
        self.notes.setText("  ".join(notes))
        self.notes.setVisible(bool(notes))
        self.status.setText(
            "Edit the text here. Click a paragraph to see where it came from."
        )

    @property
    def page(self) -> PageText:
        return self._document.pages[self._page_index]

    # -- interaction -----------------------------------------------------

    def _on_cursor(self) -> None:
        if self._loading or not self._document or not self._block_starts:
            return
        position = self.editor.textCursor().position()
        chosen = None
        for block_index, start in self._block_starts:
            if start <= position:
                chosen = block_index
            else:
                break
        if chosen is None:
            return
        self.image.highlight(self.page.blocks[chosen].bbox)

    def _on_edited(self) -> None:
        if not self._loading:
            self.status.setText("Edited. Your changes are kept when you change page.")

    def _save_edits(self) -> None:
        """Write the editor back into the page, as one block per paragraph.

        Editing collapses the page's blocks into paragraphs, which loses the
        per block boxes for anything the user restructured. That is the right
        trade: the text is what they are keeping.
        """
        if self._document is None or self._loading:
            return
        page = self._document.pages[self._page_index]
        text = self.editor.toPlainText()
        if not text.strip() and any(b.text.strip() for b in page.blocks):
            # Emptying the editor is never how someone deletes a page's text,
            # and treating it as such loses work irrecoverably.
            return
        current = "\n\n".join(
            b.text.strip() for b in sorted(page.blocks, key=lambda b: b.reading_order)
            if b.text.strip()
        )
        if text == current:
            return

        from ..core.models import TextBlock
        from ..ocr.layout import direction_of

        ordered = sorted(page.blocks, key=lambda b: b.reading_order)
        paragraphs = [p for p in text.split("\n\n") if p.strip()]
        rebuilt = []
        for index, paragraph in enumerate(paragraphs):
            # Keep the original box where the paragraph count still lines up.
            bbox = ordered[index].bbox if index < len(ordered) else (0, 0, 0, 0)
            confidence = ordered[index].confidence if index < len(ordered) else 1.0
            rebuilt.append(TextBlock(
                text=paragraph.strip(), bbox=bbox,
                direction=direction_of(paragraph), confidence=confidence,
                reading_order=index,
            ))
        page.blocks = rebuilt

    # -- clipboard and export --------------------------------------------

    def copy_page(self) -> None:
        QApplication.clipboard().setText(self.editor.toPlainText())

    def copy_all(self) -> None:
        self._save_edits()
        if self._document:
            QApplication.clipboard().setText(self._document.text())

    def export(self) -> None:
        if self._document is None:
            return
        self._save_edits()
        default = self._settings.export.default_folder or str(Path.home())
        stem = Path(self._document.source_path).stem or "pages"
        path, _chosen = QFileDialog.getSaveFileName(
            self, "Export text", str(Path(default) / f"{stem}.txt"), EXPORT_FILTERS
        )
        if not path:
            return
        try:
            self.export_to(Path(path))
        except Exception as exc:
            QMessageBox.warning(self, "Could not export", str(exc))
            return
        QMessageBox.information(self, "Exported", f"Saved to {Path(path).name}.")

    def export_to(self, path: Path) -> Path:
        if self._document is None:
            raise RuntimeError("There is nothing to export.")
        self._save_edits()
        suffix = path.suffix.lower()
        stem = Path(self._document.source_path).stem

        if suffix == ".pdf":
            return pdf_export.write_document(self._document, path, title=stem)
        if suffix == ".docx":
            return _write_docx(self._document, path, title=stem)
        path.write_text(self._document.text(), encoding="utf-8-sig")
        return path


def _write_docx(document: Document, path: Path, title: str = "") -> Path:
    """Word export for a page document, with Hebrew marked as complex script."""
    try:
        from docx import Document as WordDocument
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt
    except ImportError as exc:
        from ..export.docx import DocxUnavailable

        raise DocxUnavailable(
            "Word export needs the python-docx package, which is not installed. "
            "Export to TXT or PDF instead."
        ) from exc

    from ..language.normalize import split_by_script

    word = WordDocument()
    style = word.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:cs"), "David")

    if title:
        word.add_heading(title, level=1)

    for page in document.pages:
        if page.page_number > 1:
            word.add_page_break()
        for block in sorted(page.blocks, key=lambda b: b.reading_order):
            if not block.text.strip():
                continue
            paragraph = word.add_paragraph()
            for content, hebrew in split_by_script(block.text):
                run = paragraph.add_run(content)
                if hebrew:
                    properties = run._element.get_or_add_rPr()
                    for tag in ("w:cs", "w:rtl"):
                        properties.append(properties.makeelement(qn(tag), {}))
            if block.direction == "rtl":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    word.save(str(path))
    return path
